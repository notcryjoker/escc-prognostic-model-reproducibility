#!/usr/bin/env python3
"""
生成 BMC Medical Research Methodology 投稿格式 Word 文档：
双倍行距、连续行号、Times New Roman 12pt
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

path = sys.argv[1] if len(sys.argv) > 1 else "manuscript/manuscript_EN.docx"
doc = Document(path)

# ---------- 页面设置：启用连续行号 ----------
for sec in doc.sections:
    sectPr = sec._sectPr
    # 移除已有 lnNumType
    for old in sectPr.findall(qn('w:lnNumType')):
        sectPr.remove(old)
    ln = OxmlElement('w:lnNumType')
    ln.set(qn('w:countBy'), '1')
    ln.set(qn('w:start'), '0')
    ln.set(qn('w:numRestart'), 'continuous')
    sectPr.append(ln)

# ---------- 正文样式：Times New Roman 12pt + 双倍行距 ----------
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
rpr = st.element.get_or_add_rPr()
rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = OxmlElement('w:rFonts'); rpr.append(rf)
for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
    rf.set(qn(a), 'Times New Roman')
pf = rpr.find(qn('w:pPr'))
if pf is None:
    pf = OxmlElement('w:pPr'); rpr.append(pf)
sp = OxmlElement('w:spacing')
sp.set(qn('w:line'), '480')      # 240 = 单倍, 480 = 双倍
sp.set(qn('w:lineRule'), 'auto')
pf.append(sp)

# ---------- 段落级：字体 + 双倍行距 ----------
def fmt_par(p, size=None, bold=None, italic=None, space_after=None):
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(space_after if space_after is not None else 0)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), 'Times New Roman')
        if size is not None:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if italic is not None:
            r.font.italic = italic

for p in doc.paragraphs:
    style = (p.style.name or '').lower()
    if 'title' in style:
        fmt_par(p, size=16, bold=True, space_after=12)
    elif 'heading 1' in style:
        fmt_par(p, size=14, bold=True, space_after=6)
    elif 'heading 2' in style:
        fmt_par(p, size=13, bold=True, space_after=6)
    elif 'heading 3' in style:
        fmt_par(p, size=12, bold=True, space_after=6)
    else:
        fmt_par(p, size=12, space_after=6)

# ---------- 表格 ----------
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    rpr = r._element.get_or_add_rPr()
                    rf = rpr.find(qn('w:rFonts'))
                    if rf is None:
                        rf = OxmlElement('w:rFonts'); rpr.append(rf)
                    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                        rf.set(qn(a), 'Times New Roman')

doc.save(path)
print(f"已生成 BMC 投稿格式: {path}")
print(f"段落数 {len(doc.paragraphs)} | 表格数 {len(doc.tables)}")
