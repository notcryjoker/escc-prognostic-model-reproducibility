#!/usr/bin/env python3
"""
为 pandoc 生成的 docx 设置中文字体（正文宋体 / 标题黑体），
并统一西文字体为 Times New Roman。
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

path = sys.argv[1] if len(sys.argv) > 1 else "manuscript/ESCC_预后模型论文.docx"
doc = Document(path)


def set_run_font(run, east_asia, ascii_font="Times New Roman", size=None, bold=None):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


# 段落：按样式设定
for p in doc.paragraphs:
    style = (p.style.name or "").lower()
    if style.startswith("heading") or style.startswith("title"):
        ea, sz, bd = "黑体", None, True
        if style.startswith("title"):
            sz = 16
        elif "heading 1" in style:
            sz = 15
        elif "heading 2" in style:
            sz = 13
        elif "heading 3" in style:
            sz = 12
    else:
        ea, sz, bd = "宋体", 11, None
    for r in p.runs:
        set_run_font(r, ea, size=sz, bold=bd)

# 表格：统一宋体 10pt
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, "宋体", size=10)

# 文档默认样式
st = doc.styles['Normal']
st.font.name = "Times New Roman"
st.font.size = Pt(11)
if st.element.rPr is not None:
    rf = st.element.rPr.find(qn('w:rFonts'))
    if rf is None:
        from docx.oxml import OxmlElement
        rf = OxmlElement('w:rFonts')
        st.element.rPr.append(rf)
    rf.set(qn('w:eastAsia'), "宋体")
    rf.set(qn('w:ascii'), "Times New Roman")
    rf.set(qn('w:hAnsi'), "Times New Roman")

doc.save(path)
print(f"字体已设置并保存: {path}")
print(f"段落数: {len(doc.paragraphs)} | 表格数: {len(doc.tables)}")
